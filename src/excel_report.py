import os

from openpyxl import Workbook
from openpyxl.utils import get_column_letter


def create_excel_report(df, kpis):

    os.makedirs("output/reports", exist_ok=True)

    file_path = "output/reports/sales_analysis_report.xlsx"

    workbook = Workbook()

    # KPI Summary
    sheet = workbook.active
    sheet.title = "KPI Summary"

    sheet.append(["KPI", "Value"])

    for key, value in kpis.items():
        sheet.append([key, value])

    # Sales Data
    data_sheet = workbook.create_sheet("Sales Data")

    data_sheet.append(list(df.columns))

    for row in df.itertuples(index=False):
        data_sheet.append(list(row))

    # Product Analysis
    product_sheet = workbook.create_sheet("Product Analysis")

    product_sales = (
        df.groupby("Product")["Sales"]
        .sum()
        .sort_values(ascending=False)
    )

    product_sheet.append(["Product", "Total Sales"])

    for product, sales in product_sales.items():
        product_sheet.append([product, sales])

    # Customer Analysis
    customer_sheet = workbook.create_sheet("Customer Analysis")

    customer_sales = (
        df.groupby("Customer")["Sales"]
        .sum()
        .sort_values(ascending=False)
    )

    customer_sheet.append(["Customer", "Total Sales"])

    for customer, sales in customer_sales.items():
        customer_sheet.append([customer, sales])

    # Adjust column widths
    for worksheet in workbook.worksheets:

        for column in worksheet.columns:

            max_length = 0
            column_letter = get_column_letter(column[0].column)

            for cell in column:
                if cell.value is not None:
                    max_length = max(
                        max_length,
                        len(str(cell.value))
                    )

            worksheet.column_dimensions[column_letter].width = max_length + 2

    workbook.save(file_path)

    print("Excel report generated successfully.")